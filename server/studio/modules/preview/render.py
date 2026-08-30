"""各格式 → 統一預覽結構。

設計原則：
- **後端只做「拿不到就沒辦法」的事**（解壓 Office XML、抽 PDF 文字、偵測編碼），
  純呈現（Markdown → HTML、程式碼上色、表格排序）留給前端。
- 任何一種格式轉換失敗都不該讓整個預覽 500：`warnings[]` 帶訊息，其餘欄位照給。
- 大檔（> `MAX_RENDER_BYTES`）只回 metadata + 下載連結。
"""
from __future__ import annotations

import base64
import csv
import datetime as _dt
import html
import io
import mimetypes
import re
from pathlib import Path
from typing import Any, Optional

MAX_RENDER_BYTES = 10 * 1024 * 1024          # 超過只回 metadata
TEXT_LIMIT = 400_000                          # 純文字最多回 400KB
XLSX_MAX_ROWS = 200
XLSX_MAX_COLS = 50
CSV_MAX_ROWS = 2000
PDF_MAX_PAGES = 200
IMAGE_DATA_URI_BUDGET = 6 * 1024 * 1024       # 內嵌圖片（docx/pptx）總量上限
SINGLE_IMAGE_LIMIT = 2 * 1024 * 1024

CODE_LANG = {
    ".py": "python", ".ts": "typescript", ".tsx": "tsx", ".js": "javascript", ".jsx": "jsx", ".mjs": "javascript",
    ".cjs": "javascript", ".json": "json", ".jsonl": "json", ".yaml": "yaml", ".yml": "yaml", ".sh": "bash",
    ".zsh": "bash", ".bash": "bash", ".sql": "sql", ".go": "go", ".rs": "rust", ".java": "java", ".c": "c",
    ".h": "c", ".hpp": "cpp", ".cpp": "cpp", ".css": "css", ".scss": "scss", ".less": "less", ".xml": "xml",
    ".svg": "xml", ".toml": "toml", ".ini": "ini", ".cfg": "ini", ".txt": "text", ".log": "text", ".rb": "ruby",
    ".php": "php", ".swift": "swift", ".kt": "kotlin", ".r": "r", ".env": "bash", ".lua": "lua", ".pl": "perl",
    ".vue": "xml", ".gradle": "groovy", ".dockerfile": "dockerfile", ".patch": "diff", ".diff": "diff",
}
TEXT_NAMES = {"Dockerfile", "Makefile", "SOUL.md", "LICENSE", "README", ".gitignore", ".env"}
ENCODINGS = ("utf-8", "big5hkscs", "cp950", "gb18030")


# ------------------------------------------------------------------ helpers

def detect_kind(path: Path) -> str:
    ext = path.suffix.lower()
    mime = mimetypes.guess_type(path.name)[0] or ""
    if ext in (".html", ".htm"):
        return "html"
    if ext == ".pdf":
        return "pdf"
    if ext in (".md", ".markdown", ".mdx"):
        return "markdown"
    if mime.startswith("image/") or ext in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff", ".avif"):
        return "image"
    if ext in (".csv", ".tsv"):
        return "csv"
    if ext == ".docx":
        return "docx"
    if ext == ".pptx":
        return "pptx"
    if ext in (".xlsx", ".xlsm"):
        return "xlsx"
    if ext in CODE_LANG or mime.startswith("text/") or path.name in TEXT_NAMES:
        return "code"
    return "binary"


def language_for(path: Path) -> str:
    return CODE_LANG.get(path.suffix.lower(), "text")


def decode_bytes(data: bytes) -> tuple[str, str]:
    """回 (文字, 偵測到的編碼)。含 UTF-8 BOM、Big5/CP950、GB18030；最後退回 latin-1。"""
    if data.startswith(b"\xef\xbb\xbf"):
        try:
            return data.decode("utf-8-sig"), "utf-8-sig"
        except UnicodeDecodeError:
            pass
    for enc in ENCODINGS:
        try:
            return data.decode(enc), enc
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("latin-1", errors="replace"), "latin-1"


def read_text(path: Path, limit: int = TEXT_LIMIT) -> tuple[str, str, bool]:
    raw = path.read_bytes()
    truncated = len(raw) > limit
    text, enc = decode_bytes(raw[:limit] if truncated else raw)
    return text, enc, truncated


SAFE_IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/webp", "image/bmp", "image/tiff"}


def _data_uri(blob: bytes, content_type: str) -> str:
    ct = (content_type or "").split(";")[0].strip().lower()
    if ct not in SAFE_IMAGE_TYPES:
        ct = "image/png" if blob[:8] == b"\x89PNG\r\n\x1a\n" else "application/octet-stream"
    return f"data:{ct};base64,{base64.b64encode(blob).decode('ascii')}"


class _ImageBudget:
    def __init__(self, budget: int = IMAGE_DATA_URI_BUDGET):
        self.left = budget
        self.skipped = 0

    def take(self, blob: bytes, content_type: str) -> Optional[str]:
        if len(blob) > SINGLE_IMAGE_LIMIT or len(blob) > self.left:
            self.skipped += 1
            return None
        uri = _data_uri(blob, content_type or "image/png")
        if uri.startswith("data:application/octet-stream"):   # 認不出來的就不要塞進 HTML
            self.skipped += 1
            return None
        self.left -= len(blob)
        return uri


# -------------------------------------------------------------------- DOCX

SAFE_LINK_SCHEMES = ("http://", "https://", "mailto:", "tel:", "ftp://")


def _safe_href(href: str) -> str:
    """docx 的超連結目標可能是 `javascript:`／`data:`；只放行安全 scheme。"""
    h = (href or "").strip()
    low = h.lower().replace("\t", "").replace("\n", "").replace("\r", "")
    if low.startswith(SAFE_LINK_SCHEMES):
        return h
    if h and not re.match(r"^[a-zA-Z][a-zA-Z0-9+.\-]*:", h):   # 相對路徑
        return h
    return ""


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _run_html(run_el, part, budget: _ImageBudget) -> str:
    """一個 <w:r>：文字（含粗／斜／底線／刪除線／上下標）＋內嵌圖片。"""
    rpr = run_el.find(f"{W}rPr")
    text = "".join(t.text or "" for t in run_el.iter(f"{W}t"))
    if run_el.find(f"{W}br") is not None and not text:
        return "<br/>"
    if run_el.find(f"{W}tab") is not None and not text:
        return " "
    out = ""
    if text:
        s = html.escape(text)
        if rpr is not None:
            va = rpr.find(f"{W}vertAlign")
            if va is not None:
                val = va.get(f"{W}val")
                if val == "superscript":
                    s = f"<sup>{s}</sup>"
                elif val == "subscript":
                    s = f"<sub>{s}</sub>"
            if _on(rpr, f"{W}strike") or _on(rpr, f"{W}dstrike"):
                s = f"<s>{s}</s>"
            if _on(rpr, f"{W}u"):
                s = f"<u>{s}</u>"
            if _on(rpr, f"{W}i"):
                s = f"<em>{s}</em>"
            if _on(rpr, f"{W}b"):
                s = f"<strong>{s}</strong>"
        out += s
    for blip in run_el.iter(f"{A}blip"):
        rid = blip.get(f"{R}embed")
        if not rid:
            continue
        try:
            image_part = part.related_parts[rid]
            uri = budget.take(image_part.blob, getattr(image_part, "content_type", "image/png"))
        except Exception:
            uri = None
        out += f'<img src="{uri}" alt=""/>' if uri else '<span class="img-skipped">［圖片］</span>'
    return out


def _on(rpr, tag: str) -> bool:
    el = rpr.find(tag)
    if el is None:
        return False
    val = el.get(f"{W}val")
    return val not in ("0", "false", "none")


def _para_inner(para, budget: _ImageBudget) -> str:
    part = para.part
    parts: list[str] = []
    for child in para._p.iterchildren():
        tag = child.tag
        if tag == f"{W}r":
            parts.append(_run_html(child, part, budget))
        elif tag == f"{W}hyperlink":
            inner = "".join(_run_html(r, part, budget) for r in child.findall(f"{W}r"))
            rid = child.get(f"{R}id")
            href = ""
            if rid:
                try:
                    rel = part.rels[rid]
                    href = _safe_href(rel.target_ref) if rel.is_external else ""
                except Exception:
                    href = ""
            anchor = child.get(f"{W}anchor")
            if href:
                parts.append(f'<a href="{html.escape(href, quote=True)}" target="_blank" rel="noreferrer">{inner or html.escape(href)}</a>')
            elif anchor:
                parts.append(f'<a href="#{html.escape(anchor, quote=True)}">{inner}</a>')
            else:
                parts.append(inner)
    return "".join(parts)


def _list_info(para) -> Optional[tuple[int, bool]]:
    """(縮排層級, 是否有序)；不是清單回 None。"""
    ppr = para._p.find(f"{W}pPr")
    numpr = ppr.find(f"{W}numPr") if ppr is not None else None
    style = (para.style.name if para.style is not None else "") or ""
    if numpr is None and "List" not in style:
        return None
    level = 0
    if numpr is not None:
        ilvl = numpr.find(f"{W}ilvl")
        if ilvl is not None:
            try:
                level = int(ilvl.get(f"{W}val") or 0)
            except ValueError:
                level = 0
    ordered = "Number" in style
    return level, ordered


def _cell_html(cell, budget: _ImageBudget) -> str:
    bits = [_para_inner(p, budget) for p in cell.paragraphs]
    return "<br/>".join(b for b in bits if b.strip()) or ""


def docx_to_html(path: Path, warnings: list[str]) -> tuple[str, dict[str, Any]]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(str(path))
    budget = _ImageBudget()
    out: list[str] = []
    open_lists: list[str] = []          # 目前開著的 <ul>/<ol> 疊層
    headings: list[dict[str, Any]] = []
    counts = {"paragraphs": 0, "tables": 0, "images": 0, "headings": 0}

    def close_lists(to: int = 0) -> None:
        while len(open_lists) > to:
            out.append(f"</{open_lists.pop()}>")

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            inner = _para_inner(para, budget)
            if not inner.strip():
                continue
            counts["paragraphs"] += 1
            counts["images"] += inner.count("<img ")
            style = (para.style.name if para.style is not None else "") or ""
            m = re.match(r"Heading (\d)", style) or re.match(r"標題 ?(\d)", style)
            li = _list_info(para)
            if m:
                close_lists()
                lvl = min(int(m.group(1)), 6)
                anchor = f"h-{len(headings)}"
                headings.append({"level": lvl, "text": re.sub(r"<[^>]+>", "", inner), "id": anchor})
                counts["headings"] += 1
                out.append(f'<h{lvl} id="{anchor}">{inner}</h{lvl}>')
            elif li is not None:
                level, ordered = li
                want = level + 1
                tagname = "ol" if ordered else "ul"
                close_lists(want)
                if len(open_lists) == want and open_lists[-1] != tagname:
                    close_lists(want - 1)
                while len(open_lists) < want:
                    open_lists.append(tagname)
                    out.append(f"<{tagname}>")
                out.append(f"<li>{inner}</li>")
            elif style in ("Title", "標題"):
                close_lists()
                anchor = f"h-{len(headings)}"
                headings.append({"level": 1, "text": re.sub(r"<[^>]+>", "", inner), "id": anchor})
                out.append(f'<h1 id="{anchor}">{inner}</h1>')
            elif style in ("Quote", "Intense Quote"):
                close_lists()
                out.append(f"<blockquote>{inner}</blockquote>")
            else:
                close_lists()
                out.append(f"<p>{inner}</p>")
        elif tag == "tbl":
            close_lists()
            table = Table(child, doc)
            counts["tables"] += 1
            rows = list(table.rows)
            body: list[str] = []
            for i, r in enumerate(rows):
                cells = "".join(
                    f"<{'th' if i == 0 else 'td'}>{_cell_html(c, budget)}</{'th' if i == 0 else 'td'}>"
                    for c in r.cells
                )
                body.append(f"<tr>{cells}</tr>")
            head = f"<thead>{body[0]}</thead>" if body else ""
            rest = "".join(body[1:])
            out.append(f"<table>{head}<tbody>{rest}</tbody></table>")
    close_lists()
    if budget.skipped:
        warnings.append(f"{budget.skipped} 張內嵌圖片太大，未內嵌")
    props = doc.core_properties
    meta = {
        "headings": headings,
        "counts": counts,
        "author": props.author or "",
        "doc_title": props.title or "",
        "created": props.created.isoformat() if props.created else "",
        "modified": props.modified.isoformat() if props.modified else "",
    }
    return "\n".join(out), meta


# -------------------------------------------------------------------- XLSX

def _fmt_cell(value: Any, number_format: str) -> str:
    from openpyxl.styles.numbers import is_date_format

    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        if isinstance(value, _dt.time):
            return value.strftime("%H:%M:%S")
        if isinstance(value, _dt.datetime):
            if value.hour or value.minute or value.second:
                return value.strftime("%Y-%m-%d %H:%M:%S")
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        fmt = (number_format or "General").split(";")[0]
        try:
            if is_date_format(number_format or ""):
                from openpyxl.utils.datetime import from_excel
                d = from_excel(value)
                if isinstance(d, (_dt.datetime, _dt.date)):
                    return d.strftime("%Y-%m-%d %H:%M:%S" if getattr(d, "hour", 0) or getattr(d, "minute", 0) else "%Y-%m-%d")
        except Exception:
            pass
        percent = "%" in fmt
        n = value * 100 if percent else value
        dec = 0
        if "." in fmt:
            tail = fmt.split(".", 1)[1]
            dec = len(re.match(r"[0#]*", tail).group(0))
        elif isinstance(value, float) and not float(value).is_integer():
            dec = 2
        thousands = "#,##" in fmt or ("," in fmt.split(".")[0] and "0" in fmt)
        try:
            s = f"{n:,.{dec}f}" if thousands else f"{n:.{dec}f}"
        except (ValueError, TypeError):
            s = str(n)
        if dec == 0 and isinstance(value, float) and not thousands:
            s = f"{int(round(n))}"
        if percent:
            s += "%"
        prefix = "$" if "$" in fmt else ("NT$" if "NT$" in fmt else "")
        return prefix + s
    return str(value)


def _freeze(ws) -> dict[str, int]:
    from openpyxl.utils import coordinate_to_tuple

    fp = getattr(ws, "freeze_panes", None)
    if not fp or not isinstance(fp, str):
        return {"rows": 0, "cols": 0}
    try:
        row, col = coordinate_to_tuple(fp)
        return {"rows": max(0, row - 1), "cols": max(0, col - 1)}
    except Exception:
        return {"rows": 0, "cols": 0}


def xlsx_to_sheets(path: Path, warnings: list[str]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    formulas = None
    try:
        formulas = openpyxl.load_workbook(str(path), data_only=False)
    except Exception:
        formulas = None
    sheets: list[dict[str, Any]] = []
    stale_formula = False
    for ws in wb.worksheets:
        fws = formulas[ws.title] if formulas is not None and ws.title in formulas.sheetnames else None
        rows: list[list[str]] = []
        numeric_hits: dict[int, int] = {}
        total_rows = ws.max_row or 0
        total_cols = ws.max_column or 0
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(total_rows, XLSX_MAX_ROWS), max_col=min(total_cols, XLSX_MAX_COLS) or 1)):
            cells: list[str] = []
            for j, c in enumerate(row):
                v = c.value
                if v is None and fws is not None:
                    raw = fws.cell(row=c.row, column=c.column).value
                    if isinstance(raw, str) and raw.startswith("="):
                        stale_formula = True
                        cells.append(raw)
                        continue
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    numeric_hits[j] = numeric_hits.get(j, 0) + 1
                cells.append(_fmt_cell(v, c.number_format))
            rows.append(cells)
            if i + 1 >= XLSX_MAX_ROWS:
                break
        body_rows = max(1, len(rows) - 1)
        numeric_cols = [j for j, n in numeric_hits.items() if n >= body_rows * 0.6]
        merges = []
        for rng in getattr(ws, "merged_cells", []).ranges if hasattr(ws, "merged_cells") else []:
            if rng.min_row > XLSX_MAX_ROWS or rng.min_col > XLSX_MAX_COLS:
                continue
            if rng.max_row == rng.min_row and rng.max_col == rng.min_col:
                continue
            merges.append({
                "row": rng.min_row - 1, "col": rng.min_col - 1,
                "rowspan": min(rng.max_row, XLSX_MAX_ROWS) - rng.min_row + 1,
                "colspan": min(rng.max_col, XLSX_MAX_COLS) - rng.min_col + 1,
            })
        sheets.append({
            "name": ws.title,
            "rows": rows,
            "merges": merges,
            "freeze": _freeze(ws),
            "numeric_cols": sorted(numeric_cols),
            "total_rows": total_rows,
            "total_cols": total_cols,
            "truncated": total_rows > XLSX_MAX_ROWS or total_cols > XLSX_MAX_COLS,
            "hidden": ws.sheet_state != "visible",
        })
    wb.close()
    if formulas is not None:
        formulas.close()
    if any(s["truncated"] for s in sheets):
        warnings.append(f"表格已截斷，只顯示前 {XLSX_MAX_ROWS} 列 × {XLSX_MAX_COLS} 欄（可下載原檔看全部）")
    if stale_formula:
        warnings.append("有公式沒有快取計算值（檔案未被 Excel 開啟過），顯示公式原文")
    return sheets, {"sheet_names": [s["name"] for s in sheets]}


# -------------------------------------------------------------------- PPTX

def pptx_to_pages(path: Path, warnings: list[str]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    budget = _ImageBudget()
    pages: list[dict[str, Any]] = []
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip()
        except Exception:
            title = ""
        body: list[str] = []
        images: list[str] = []
        tables: list[list[list[str]]] = []
        for shape in slide.shapes:
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    uri = budget.take(shape.image.blob, shape.image.content_type)
                    if uri:
                        images.append(uri)
                    continue
            except Exception:
                pass
            if getattr(shape, "has_table", False) and shape.has_table:
                tables.append([[c.text for c in r.cells] for r in shape.table.rows])
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t and t != title:
                        body.append(("  " * max(0, para.level or 0)) + t)
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        except Exception:
            notes = ""
        pages.append({"index": i, "title": title, "body": body, "notes": notes, "images": images, "tables": tables})
    if budget.skipped:
        warnings.append(f"{budget.skipped} 張投影片圖片太大，未內嵌")
    return pages, {"slides": len(pages), "size": {"w": prs.slide_width, "h": prs.slide_height}}


# --------------------------------------------------------------------- PDF

def pdf_to_pages(path: Path, warnings: list[str]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - 相依沒裝時只失去文字，仍能用 <embed> 看原檔
        warnings.append("未安裝 pypdf，無法抽出 PDF 文字")
        return [], {"pages": 0}
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            warnings.append("PDF 已加密，無法抽出文字")
            return [], {"pages": len(reader.pages), "encrypted": True}
    pages: list[dict[str, Any]] = []
    for i, page in enumerate(reader.pages[:PDF_MAX_PAGES], 1):
        try:
            text = page.extract_text() or ""
        except Exception as e:
            text = ""
            warnings.append(f"第 {i} 頁文字抽取失敗：{e}")
        pages.append({"index": i, "text": text})
    total = len(reader.pages)
    if total > PDF_MAX_PAGES:
        warnings.append(f"只抽出前 {PDF_MAX_PAGES} 頁的文字（共 {total} 頁）")
    info = {}
    try:
        md = reader.metadata or {}
        info = {"doc_title": str(md.get("/Title") or ""), "author": str(md.get("/Author") or "")}
    except Exception:
        info = {}
    return pages, {"pages": total, **info}


# --------------------------------------------------------------------- CSV

def csv_to_rows(path: Path, warnings: list[str]) -> tuple[list[list[str]], dict[str, Any]]:
    raw = path.read_bytes()
    text, encoding = decode_bytes(raw)
    sample = text[:8192]
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        pass
    rows: list[list[str]] = []
    truncated = False
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    try:
        for i, r in enumerate(reader):
            if i >= CSV_MAX_ROWS:
                truncated = True
                break
            rows.append(r)
    except csv.Error as e:
        warnings.append(f"CSV 解析中斷：{e}")
    if truncated:
        warnings.append(f"只顯示前 {CSV_MAX_ROWS} 列")
    width = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (width - len(r)) for r in rows]
    numeric_cols = []
    body = rows[1:] if len(rows) > 1 else []
    for j in range(width):
        vals = [r[j] for r in body if r[j].strip()]
        if vals and sum(1 for v in vals if _isnum(v)) >= len(vals) * 0.8:
            numeric_cols.append(j)
    return rows, {"encoding": encoding, "delimiter": delimiter, "numeric_cols": numeric_cols,
                  "total_rows": len(rows), "total_cols": width, "truncated": truncated}


def _isnum(s: str) -> bool:
    try:
        float(s.replace(",", "").replace("%", "").strip())
        return True
    except ValueError:
        return False


# ------------------------------------------------------------------- IMAGE

_EXIF_KEYS = ("Make", "Model", "DateTime", "DateTimeOriginal", "ExposureTime", "FNumber", "ISOSpeedRatings",
              "FocalLength", "LensModel", "Orientation", "Software")


def image_meta(path: Path, warnings: list[str]) -> dict[str, Any]:
    try:
        from PIL import Image, ExifTags
    except ImportError:  # pragma: no cover
        return {}
    try:
        with Image.open(str(path)) as im:
            meta: dict[str, Any] = {"width": im.width, "height": im.height, "format": im.format or "",
                                    "mode": im.mode, "animated": bool(getattr(im, "is_animated", False))}
            exif_out: dict[str, str] = {}
            try:
                exif = im.getexif()
                tags = {v: k for k, v in ExifTags.TAGS.items()}
                for key in _EXIF_KEYS:
                    tid = tags.get(key)
                    if tid is None:
                        continue
                    val = exif.get(tid)
                    if val in (None, ""):
                        continue
                    exif_out[key] = str(val)
                if exif.get_ifd(0x8825):
                    exif_out["GPS"] = "有座標資料"
            except Exception:
                pass
            if exif_out:
                meta["exif"] = exif_out
            return meta
    except Exception as e:
        warnings.append(f"讀取圖片資訊失敗：{e}")
        return {}


# ------------------------------------------------------------------ 主入口

def render(path: Path, kind: str = "auto") -> dict[str, Any]:
    """把一個實體檔案轉成統一預覽結構（不含 url／權限，那是 router 的事）。"""
    st = path.stat()
    if kind in ("", "auto", None):
        kind = detect_kind(path)
    warnings: list[str] = []
    out: dict[str, Any] = {
        "kind": kind,
        "title": path.name,
        "path": str(path),
        "size": st.st_size,
        "mtime": st.st_mtime,
        "mime": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
        "meta": {},
        "warnings": warnings,
    }
    if st.st_size > MAX_RENDER_BYTES:
        out["too_large"] = True
        warnings.append(f"檔案 {st.st_size // 1024 // 1024} MB 超過 10MB，只顯示基本資訊，請下載原檔")
        if kind in ("image", "pdf"):
            out["meta"] = image_meta(path, warnings) if kind == "image" else {}
        return out
    try:
        if kind in ("markdown", "code", "html"):
            text, encoding, truncated = read_text(path)
            out["text"] = text
            out["meta"] = {"encoding": encoding, "language": language_for(path), "lines": text.count("\n") + 1}
            if truncated:
                warnings.append(f"檔案過大，只顯示前 {TEXT_LIMIT // 1000}KB")
        elif kind == "csv":
            rows, meta = csv_to_rows(path, warnings)
            out["rows"] = rows
            out["meta"] = meta
        elif kind == "docx":
            html_out, meta = docx_to_html(path, warnings)
            out["html"] = html_out
            out["meta"] = meta
        elif kind == "xlsx":
            sheets, meta = xlsx_to_sheets(path, warnings)
            out["sheets"] = sheets
            out["meta"] = meta
        elif kind == "pptx":
            pages, meta = pptx_to_pages(path, warnings)
            out["pages"] = pages
            out["meta"] = meta
        elif kind == "pdf":
            pages, meta = pdf_to_pages(path, warnings)
            out["pages"] = pages
            out["meta"] = meta
        elif kind == "image":
            out["meta"] = image_meta(path, warnings)
        # binary：只有 metadata + 下載連結
    except Exception as e:
        out["kind"] = "binary"
        out["error"] = f"預覽轉換失敗：{type(e).__name__}: {e}"
        warnings.append(out["error"])
    return out


def render_inline(text: str, kind: str = "auto", title: str = "", language: str = "") -> dict[str, Any]:
    """非檔案的純文字內容（工作流節點輸出、審批 payload…）。"""
    warnings: list[str] = []
    body = text or ""
    if len(body) > TEXT_LIMIT:
        body = body[:TEXT_LIMIT]
        warnings.append(f"內容過大，只顯示前 {TEXT_LIMIT // 1000}KB")
    if kind in ("", "auto", None):
        kind = "markdown"
    out: dict[str, Any] = {
        "kind": kind,
        "title": title or "output",
        "path": "",
        "size": len((text or "").encode("utf-8")),
        "mtime": 0.0,
        "mime": "text/markdown" if kind == "markdown" else "text/plain",
        "meta": {"language": language or "text", "lines": body.count("\n") + 1, "inline": True},
        "warnings": warnings,
    }
    if kind == "csv":
        reader = csv.reader(io.StringIO(body))
        rows = [r for _, r in zip(range(CSV_MAX_ROWS), reader)]
        width = max((len(r) for r in rows), default=0)
        out["rows"] = [r + [""] * (width - len(r)) for r in rows]
        out["meta"].update(total_rows=len(rows), total_cols=width, numeric_cols=[], delimiter=",", encoding="utf-8")
    else:
        out["text"] = body
    return out


__all__ = ["render", "render_inline", "detect_kind", "decode_bytes", "MAX_RENDER_BYTES",
           "XLSX_MAX_ROWS", "XLSX_MAX_COLS"]
