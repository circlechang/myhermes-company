"""上傳／下載／預覽。安全邊界：

- 上傳只寫到 `<db 目錄>/uploads/<company_id>/<session_id>/`，檔名淨化、同名加序號。
- 讀取（下載／預覽）只允許三種路徑：
  1. 該公司的 uploads 目錄；
  2. Hermes workspace（`$HERMES_HOME/workspace`）；
  3. 這位成員自己 session 訊息內文／附件裡「原文出現過」的絕對路徑（agent 產出的檔案）。
  其餘一律 404（不洩漏是否存在）。symlink 先 resolve 再比對。
"""
from __future__ import annotations

import csv
import html
import io
import json
import mimetypes
import re
from pathlib import Path
from typing import Any, Optional

from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import FileResponse
from sqlmodel import Session, select

from ...auth import Principal, _resolve, current_principal, get_db
from ...errors import ApiError, bad_request, not_found
from ...models import ChatSession, Message

router = APIRouter()

MAX_UPLOAD_BYTES = 50 * 1024 * 1024
TEXT_PREVIEW_LIMIT = 400_000
CODE_EXT = {
    ".py": "python", ".ts": "typescript", ".tsx": "tsx", ".js": "javascript", ".jsx": "jsx", ".json": "json",
    ".yaml": "yaml", ".yml": "yaml", ".sh": "bash", ".zsh": "bash", ".sql": "sql", ".go": "go", ".rs": "rust",
    ".java": "java", ".c": "c", ".h": "c", ".cpp": "cpp", ".css": "css", ".scss": "scss", ".xml": "xml",
    ".toml": "toml", ".ini": "ini", ".txt": "text", ".log": "text", ".rb": "ruby", ".php": "php", ".swift": "swift",
    ".kt": "kotlin", ".r": "r", ".env": "bash",
}


def principal_header_or_query(request: Request, db: Session = Depends(get_db)) -> Principal:
    """<img>/<iframe> 帶不了 Authorization header，所以檔案讀取也接受 `?token=`。"""
    auth = request.headers.get("authorization", "")
    if auth.lower().startswith("bearer "):
        return current_principal(request, db)
    token = request.query_params.get("token") or ""
    if not token:
        raise ApiError(401, "unauthorized", "missing bearer token")
    return _resolve(token, request.app.state.settings.secret, db)


def upload_root(settings) -> Path:
    return Path(settings.db_path).expanduser().parent / "uploads"


def ensure_upload_root(settings) -> None:
    try:
        upload_root(settings).mkdir(parents=True, exist_ok=True)
    except OSError:
        pass


def workspace_root(settings) -> Path:
    return Path(settings.hermes_home).expanduser() / "workspace"


_SAFE_NAME = re.compile(r"[^\w.\-一-鿿぀-ヿ ]+")


def safe_filename(name: str) -> str:
    base = Path(name or "file").name
    base = _SAFE_NAME.sub("_", base).strip(" .") or "file"
    return base[:120]


def unique_path(d: Path, name: str) -> Path:
    p = d / name
    if not p.exists():
        return p
    stem, suffix = Path(name).stem, Path(name).suffix
    for i in range(1, 1000):
        q = d / f"{stem}-{i}{suffix}"
        if not q.exists():
            return q
    raise bad_request("too many files with the same name")


def _owned_session(db: Session, p: Principal, sid: str) -> ChatSession:
    s = db.get(ChatSession, sid)
    if s is None or s.company_id != p.company_id or (s.member_id != p.member.id and p.role not in ("owner", "admin")):
        raise not_found("session")
    return s


def file_public(path: Path, session_id: str) -> dict[str, Any]:
    st = path.stat()
    return {"name": path.name, "path": str(path), "size": st.st_size,
            "mime": mimetypes.guess_type(path.name)[0] or "application/octet-stream", "session_id": session_id}


@router.post("/uploads", status_code=201)
async def upload(request: Request, session_id: str = Form(...), files: list[UploadFile] = File(...),
                 p: Principal = Depends(current_principal), db: Session = Depends(get_db)):
    s = _owned_session(db, p, session_id)
    d = upload_root(request.app.state.settings) / p.company_id / s.id
    d.mkdir(parents=True, exist_ok=True)
    out = []
    for f in files:
        target = unique_path(d, safe_filename(f.filename or "file"))
        size = 0
        with target.open("wb") as fh:
            while True:
                chunk = await f.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > MAX_UPLOAD_BYTES:
                    fh.close()
                    target.unlink(missing_ok=True)
                    raise ApiError(413, "too_large", f"{f.filename} 超過 50MB")
                fh.write(chunk)
        out.append(file_public(target, s.id))
    return out


@router.get("/uploads")
def list_uploads(request: Request, session_id: str, p: Principal = Depends(current_principal), db: Session = Depends(get_db)):
    s = _owned_session(db, p, session_id)
    d = upload_root(request.app.state.settings) / p.company_id / s.id
    if not d.exists():
        return []
    return [file_public(x, s.id) for x in sorted(d.iterdir()) if x.is_file()]


# ---- read-side whitelist -------------------------------------------------

def _is_under(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False


def _mentioned_in_messages(db: Session, p: Principal, raw: str) -> bool:
    """Path string appears verbatim in a message (content or attachments) of a session the member can read."""
    q = (select(Message.id).join(ChatSession, ChatSession.id == Message.session_id)
         .where(ChatSession.company_id == p.company_id))
    if p.role not in ("owner", "admin"):
        q = q.where(ChatSession.member_id == p.member.id)
    like = f"%{raw}%"
    q = q.where((Message.content.like(like)) | (Message.attachments.like(like)) | (Message.tool_result.like(like))).limit(1)
    return db.exec(q).first() is not None


def resolve_readable(request: Request, db: Session, p: Principal, raw: str) -> Path:
    raw = (raw or "").strip()
    if not raw or not raw.startswith(("/", "~")):
        raise not_found("file")
    path = Path(raw).expanduser()
    try:
        real = path.resolve(strict=True)
    except (OSError, RuntimeError):
        raise not_found("file")
    if not real.is_file():
        raise not_found("file")
    settings = request.app.state.settings
    roots = [upload_root(settings).resolve() / p.company_id, workspace_root(settings).resolve()]
    if any(_is_under(real, r) for r in roots):
        return real
    if _mentioned_in_messages(db, p, raw) or (str(real) != raw and _mentioned_in_messages(db, p, str(real))):
        return real
    raise not_found("file")


@router.get("/files")
def download(request: Request, path: str, download: bool = False, p: Principal = Depends(principal_header_or_query),
             db: Session = Depends(get_db)):
    real = resolve_readable(request, db, p, path)
    mime = mimetypes.guess_type(real.name)[0] or "application/octet-stream"
    headers = {}
    if download:
        headers["Content-Disposition"] = f"attachment; filename*=UTF-8''{_rfc5987(real.name)}"
    else:
        headers["Content-Disposition"] = f"inline; filename*=UTF-8''{_rfc5987(real.name)}"
    if mime == "text/html":
        # never execute agent-produced HTML in the app origin; the UI renders it in a sandboxed iframe via srcdoc
        mime = "text/plain; charset=utf-8"
    return FileResponse(str(real), media_type=mime, headers=headers)


def _rfc5987(name: str) -> str:
    from urllib.parse import quote
    return quote(name, safe="")


# ---- structured preview --------------------------------------------------

def preview_kind(path: Path) -> str:
    ext = path.suffix.lower()
    mime = mimetypes.guess_type(path.name)[0] or ""
    if ext in (".html", ".htm"):
        return "html"
    if ext == ".pdf":
        return "pdf"
    if mime.startswith("image/"):
        return "image"
    if ext in (".md", ".markdown"):
        return "markdown"
    if ext in (".csv", ".tsv"):
        return "csv"
    if ext == ".docx":
        return "docx"
    if ext == ".pptx":
        return "pptx"
    if ext in (".xlsx", ".xlsm"):
        return "xlsx"
    if ext in CODE_EXT or mime.startswith("text/"):
        return "code"
    return "binary"


def _read_text(path: Path) -> tuple[str, bool]:
    data = path.read_bytes()[: TEXT_PREVIEW_LIMIT + 1]
    truncated = len(data) > TEXT_PREVIEW_LIMIT
    return data[:TEXT_PREVIEW_LIMIT].decode("utf-8", errors="replace"), truncated


def docx_to_html(path: Path) -> str:
    import docx  # python-docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(str(path))
    out: list[str] = []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            style = (para.style.name if para.style is not None else "") or ""
            txt = html.escape(para.text)
            if not txt.strip():
                continue
            m = re.match(r"Heading (\d)", style)
            if m:
                lvl = min(int(m.group(1)), 6)
                out.append(f"<h{lvl}>{txt}</h{lvl}>")
            elif "List" in style:
                out.append(f"<li>{txt}</li>")
            else:
                out.append(f"<p>{txt}</p>")
        elif tag == "tbl":
            table = Table(child, doc)
            rows = []
            for r in table.rows:
                cells = "".join(f"<td>{html.escape(c.text)}</td>" for c in r.cells)
                rows.append(f"<tr>{cells}</tr>")
            out.append("<table>" + "".join(rows) + "</table>")
    return "\n".join(out)


def pptx_to_html(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        parts = [f"<section class='slide'><h3>Slide {i}</h3>"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        parts.append(f"<p>{html.escape(t)}</p>")
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = []
                for r in shape.table.rows:
                    rows.append("<tr>" + "".join(f"<td>{html.escape(c.text)}</td>" for c in r.cells) + "</tr>")
                parts.append("<table>" + "".join(rows) + "</table>")
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if notes:
            parts.append(f"<p class='notes'><em>{html.escape(notes)}</em></p>")
        parts.append("</section>")
        out.append("".join(parts))
    return "\n".join(out)


def xlsx_to_sheets(path: Path, max_rows: int = 500, max_cols: int = 50) -> list[dict[str, Any]]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows: list[list[Any]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            rows.append([("" if v is None else v) for v in row[:max_cols]])
        sheets.append({"name": ws.title, "rows": rows, "truncated": (ws.max_row or 0) > max_rows})
    wb.close()
    return sheets


def csv_rows(path: Path, max_rows: int = 1000) -> dict[str, Any]:
    text, _ = _read_text(path)
    dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = []
    for i, r in enumerate(reader):
        if i >= max_rows:
            return {"rows": rows, "truncated": True}
        rows.append(r)
    return {"rows": rows, "truncated": False}


@router.get("/preview")
def preview(request: Request, path: str, p: Principal = Depends(current_principal), db: Session = Depends(get_db)):
    real = resolve_readable(request, db, p, path)
    kind = preview_kind(real)
    base: dict[str, Any] = {"kind": kind, "name": real.name, "path": str(real), "size": real.stat().st_size,
                            "url": f"/chat/files?path={_rfc5987(str(real))}"}
    try:
        if kind in ("html", "markdown", "code"):
            text, truncated = _read_text(real)
            base.update(text=text, truncated=truncated, language=CODE_EXT.get(real.suffix.lower(), "text"))
        elif kind == "csv":
            base.update(csv_rows(real))
        elif kind == "docx":
            base.update(html=docx_to_html(real))
        elif kind == "pptx":
            base.update(html=pptx_to_html(real))
        elif kind == "xlsx":
            base.update(sheets=xlsx_to_sheets(real))
        # pdf / image / binary: client uses `url`
    except Exception as e:  # conversion failure should not 500 the UI
        base.update(kind="binary", error=f"預覽轉換失敗：{e}")
    return base


__all__ = ["router", "ensure_upload_root", "upload_root", "resolve_readable", "preview_kind", "json"]
