"""統一預覽模組 `/preview`：各格式轉換、路徑白名單／防穿越、大檔保護、編碼偵測、inline。"""
from __future__ import annotations

import json
from pathlib import Path

import pytest

from preview_samples import (
    make_code,
    make_csv_big5,
    make_csv_utf8_bom,
    make_docx,
    make_markdown,
    make_pdf,
    make_png,
    make_pptx,
    make_tsv,
    make_xlsx,
)


def _agent(client, auth, profile="default"):
    return next(a["id"] for a in client.get("/agents", headers=auth).json() if a["profile"] == profile)


def _session(client, auth, **kw):
    return client.post("/sessions", json={"agent_id": _agent(client, auth), **kw}, headers=auth).json()


def _in_workspace(app, name: str) -> Path:
    ws = Path(app.state.settings.hermes_home) / "workspace"
    ws.mkdir(parents=True, exist_ok=True)
    return ws / name


def _preview(client, auth, path: str, **params):
    return client.get("/preview", params={"path": str(path), **params}, headers=auth)


# ------------------------------------------------------------- 路徑與權限

def test_absolute_path_whitelist_and_traversal(client, auth, app, tmp_path):
    """workspace 內可讀；workspace 外、`..` 穿越、相對路徑一律 404/400。"""
    good = make_markdown(_in_workspace(app, "ok.md"))
    assert _preview(client, auth, good).status_code == 200

    secret = tmp_path / "secret.txt"
    secret.write_text("nope")
    assert _preview(client, auth, secret).status_code == 404

    ws = Path(app.state.settings.hermes_home) / "workspace"
    assert _preview(client, auth, f"{ws}/../../secret.txt").status_code == 404
    assert _preview(client, auth, "workspace/../../../etc/passwd").status_code == 400
    assert _preview(client, auth, "workspace/..%2f..%2fetc/passwd").status_code in (400, 404)
    assert _preview(client, auth, "no-such-root/x.md").status_code == 404
    assert _preview(client, auth, "").status_code == 404


def test_symlink_escaping_the_root_is_rejected(client, auth, app, tmp_path):
    outside = tmp_path / "outside.md"
    outside.write_text("# secret")
    link = _in_workspace(app, "link.md")
    link.symlink_to(outside)
    # 絕對路徑會先 resolve 再比對根 → 逃出去就 404
    assert _preview(client, auth, link).status_code == 404
    # 虛擬路徑走檔案瀏覽器的 resolve()，同樣擋掉
    assert _preview(client, auth, "workspace/link.md").status_code in (400, 404)


def test_virtual_path_and_uploads_path(client, auth, app):
    make_markdown(_in_workspace(app, "note.md"))
    r = _preview(client, auth, "workspace/note.md")
    assert r.status_code == 200
    body = r.json()
    assert body["kind"] == "markdown" and body["title"] == "note.md"
    assert body["url"].startswith("/preview/raw/note.md?path=") and body["download_url"].endswith("&download=true")

    s = _session(client, auth)
    up = client.post("/chat/uploads", data={"session_id": s["id"]},
                     files=[("files", ("up.md", b"# uploaded"))], headers=auth).json()
    assert _preview(client, auth, up[0]["path"]).json()["text"] == "# uploaded"


def test_path_mentioned_in_message_is_readable(client, auth, token, tmp_path):
    out = tmp_path / "agent-output.csv"
    out.write_text("a,b\n1,2\n")
    s = _session(client, auth)
    with client.websocket_connect(f"/ws/chat?token={token}") as ws:
        ws.receive_text()
        ws.send_text(json.dumps({"type": "run", "session_id": s["id"], "input": f"存到 {out}"}))
        while True:
            if json.loads(ws.receive_text())["type"] in ("run.completed", "run.failed"):
                break
    assert _preview(client, auth, out).json()["rows"] == [["a", "b"], ["1", "2"]]
    (tmp_path / "never-mentioned.csv").write_text("x")
    assert _preview(client, auth, tmp_path / "never-mentioned.csv").status_code == 404


def test_preview_requires_auth(client, app):
    make_markdown(_in_workspace(app, "auth.md"))
    assert client.get("/preview", params={"path": "workspace/auth.md"}).status_code == 401


# ----------------------------------------------------------------- 各格式

def test_docx_rich_html(client, auth, app):
    make_docx(_in_workspace(app, "report.docx"))
    b = _preview(client, auth, "workspace/report.docx").json()
    assert b["kind"] == "docx"
    h = b["html"]
    assert '<h1 id="h-0">季度營運報告</h1>' in h and "<h2" in h
    assert "<strong>成長 30%</strong>" in h and "<em>循環箱</em>" in h
    assert "<ul>" in h and "<ol>" in h and "<li>第一點</li>" in h
    assert "<thead><tr><th>品項</th>" in h and "<td>循環箱</td>" in h
    assert 'src="data:image/png;base64,' in h            # 內嵌圖片轉 data URI
    assert [x["text"] for x in b["meta"]["headings"]][:2] == ["季度營運報告", "一、摘要"]
    assert b["meta"]["counts"]["tables"] == 1 and b["meta"]["counts"]["images"] == 1


def test_docx_escapes_html(client, auth, app):
    import docx

    d = docx.Document()
    d.add_paragraph("<script>alert(1)</script>")
    d.save(str(_in_workspace(app, "xss.docx")))
    h = _preview(client, auth, "workspace/xss.docx").json()["html"]
    assert "<script>" not in h and "&lt;script&gt;" in h


def test_docx_hyperlinks_are_scheme_checked(client, auth, app):
    """docx 可以塞 `javascript:` 超連結；轉出來的 HTML 不能帶著它。"""
    import docx
    from docx.oxml.ns import qn

    d = docx.Document()
    para = d.add_paragraph()
    rid_bad = d.part.relate_to("javascript:alert(1)",
                               "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                               is_external=True)
    rid_ok = d.part.relate_to("https://example.com/a",
                              "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                              is_external=True)
    for rid, text in ((rid_bad, "壞連結"), (rid_ok, "好連結")):
        h = para._p.makeelement(qn("w:hyperlink"), {qn("r:id"): rid})
        h.append(para.add_run(text)._r)
        para._p.append(h)
    d.save(str(_in_workspace(app, "links.docx")))
    h = _preview(client, auth, "workspace/links.docx").json()["html"]
    assert "javascript" not in h and "壞連結" in h
    assert '<a href="https://example.com/a"' in h and "好連結" in h


def test_xlsx_sheets_formats_merges_freeze_and_truncation(client, auth, app):
    make_xlsx(_in_workspace(app, "revenue.xlsx"), big=True)
    b = _preview(client, auth, "workspace/revenue.xlsx").json()
    assert b["kind"] == "xlsx"
    names = [s["name"] for s in b["sheets"]]
    assert names == ["營收", "備註", "大表"]
    first = b["sheets"][0]
    assert first["rows"][1][1] == "1,234,567.50"          # 千分位＋兩位小數
    assert first["rows"][1][2] == "12.8%"                  # 百分比格式
    assert first["rows"][1][3] == "2026-01-15"             # 日期格式
    assert first["freeze"] == {"rows": 1, "cols": 1}
    assert {"row": 4, "col": 0, "rowspan": 1, "colspan": 3} in first["merges"]
    assert first["numeric_cols"] == [1, 2]
    big = b["sheets"][2]
    assert len(big["rows"]) == 200 and len(big["rows"][0]) == 50 and big["truncated"] is True
    assert any("截斷" in w for w in b["warnings"])


def test_xlsx_formula_without_cached_value_shows_formula(client, auth, app):
    make_xlsx(_in_workspace(app, "f.xlsx"))
    b = _preview(client, auth, "workspace/f.xlsx").json()
    assert b["sheets"][0]["rows"][1][4] == "=B2*2"
    assert any("公式" in w for w in b["warnings"])


def test_pptx_slides_notes_and_images(client, auth, app):
    make_pptx(_in_workspace(app, "deck.pptx"), slides=3)
    b = _preview(client, auth, "workspace/deck.pptx").json()
    assert b["kind"] == "pptx" and b["meta"]["slides"] == 3 and len(b["pages"]) == 3
    p1 = b["pages"][0]
    assert p1["title"] == "第 1 頁：主題" and "重點 1-1" in p1["body"] and "備註 1" in p1["notes"]
    assert b["pages"][1]["images"] and b["pages"][1]["images"][0].startswith("data:image/png;base64,")


def test_pdf_pages_text(client, auth, app):
    make_pdf(_in_workspace(app, "doc.pdf"), ["Hello preview page one", "Second page content"])
    b = _preview(client, auth, "workspace/doc.pdf").json()
    assert b["kind"] == "pdf" and b["meta"]["pages"] == 2
    assert [p["index"] for p in b["pages"]] == [1, 2]
    assert "Hello preview" in b["pages"][0]["text"] and "Second page" in b["pages"][1]["text"]
    assert b["url"].startswith("/preview/raw/doc.pdf?")    # 前端用 <object> 顯示原檔


def test_csv_encoding_and_delimiter_detection(client, auth, app):
    make_csv_big5(_in_workspace(app, "big5.csv"))
    b = _preview(client, auth, "workspace/big5.csv").json()
    assert b["kind"] == "csv" and b["meta"]["encoding"] in ("big5hkscs", "cp950")
    assert b["rows"] == [["品項", "數量", "單價"], ["循環箱", "1200", "30"], ["提袋", "800", "12"]]
    assert b["meta"]["numeric_cols"] == [1, 2]

    make_csv_utf8_bom(_in_workspace(app, "bom.csv"))
    b = _preview(client, auth, "workspace/bom.csv").json()
    assert b["meta"]["encoding"] == "utf-8-sig" and b["rows"][0] == ["name", "score"]

    make_tsv(_in_workspace(app, "d.tsv"))
    b = _preview(client, auth, "workspace/d.tsv").json()
    assert b["meta"]["delimiter"] == "\t" and b["rows"][1] == ["1", "2", "3"]


def test_image_meta_and_code(client, auth, app):
    make_png(_in_workspace(app, "pixel.png"))
    b = _preview(client, auth, "workspace/pixel.png").json()
    assert b["kind"] == "image" and b["meta"]["width"] == 64 and b["meta"]["height"] == 48
    assert b["meta"]["format"] == "PNG"

    make_code(_in_workspace(app, "box.py"))
    b = _preview(client, auth, "workspace/box.py").json()
    assert b["kind"] == "code" and b["meta"]["language"] == "python" and b["meta"]["encoding"] == "utf-8"
    assert "class Box" in b["text"]


def test_markdown_is_returned_raw(client, auth, app):
    make_markdown(_in_workspace(app, "r.md"))
    b = _preview(client, auth, "workspace/r.md").json()
    assert b["kind"] == "markdown" and "html" not in b   # 後端不渲染 Markdown
    assert b["text"].startswith("# 循環箱專案報告")
    assert b["mtime"] > 0 and b["size"] > 0


def test_kind_override(client, auth, app):
    make_markdown(_in_workspace(app, "as-code.md"))
    b = _preview(client, auth, "workspace/as-code.md", kind="code").json()
    assert b["kind"] == "code"


# --------------------------------------------------------------- 大檔／錯誤

def test_large_file_returns_metadata_only(client, auth, app):
    big = _in_workspace(app, "huge.md")
    big.write_bytes(b"# big\n" + b"x" * (11 * 1024 * 1024))
    b = _preview(client, auth, "workspace/huge.md").json()
    assert b["too_large"] is True and "text" not in b
    assert b["size"] > 10 * 1024 * 1024 and b["download_url"]
    assert any("10MB" in w for w in b["warnings"])


def test_broken_office_file_does_not_500(client, auth, app):
    bad = _in_workspace(app, "broken.docx")
    bad.write_bytes(b"not a docx at all")
    b = _preview(client, auth, "workspace/broken.docx").json()
    assert b["kind"] == "binary" and b["error"] and b["warnings"]


# ------------------------------------------------------------------ inline

def test_inline_preview(client, auth):
    r = client.post("/preview/inline", json={"text": "# 節點輸出\n\n- a\n- b", "title": "node-1"}, headers=auth)
    assert r.status_code == 200
    b = r.json()
    assert b["kind"] == "markdown" and b["title"] == "node-1" and b["meta"]["inline"] is True
    assert b["text"].startswith("# 節點輸出") and b["size"] > 0

    b = client.post("/preview/inline", json={"text": "a,b\n1,2", "kind": "csv"}, headers=auth).json()
    assert b["rows"] == [["a", "b"], ["1", "2"]]

    b = client.post("/preview/inline", json={"text": "x" * 500_000}, headers=auth).json()
    assert len(b["text"]) == 400_000 and b["warnings"]

    assert client.post("/preview/inline", json={"text": "x"}).status_code == 401


# --------------------------------------------------------------------- raw

def test_raw_serves_bytes_with_query_token(client, auth, token, app):
    make_png(_in_workspace(app, "raw.png"))
    r = client.get("/preview/raw", params={"path": "workspace/raw.png", "token": token})
    assert r.status_code == 200 and r.headers["content-type"].startswith("image/png")
    assert r.headers["x-content-type-options"] == "nosniff"
    r = client.get("/preview/raw", params={"path": "workspace/raw.png", "download": "true"}, headers=auth)
    assert "attachment" in r.headers["content-disposition"]
    assert client.get("/preview/raw", params={"path": "workspace/raw.png"}).status_code == 401


def test_raw_never_serves_html_as_html(client, auth, app):
    _in_workspace(app, "x.html").write_bytes(b"<script>alert(1)</script>")
    r = client.get("/preview/raw", params={"path": "workspace/x.html"}, headers=auth)
    assert r.headers["content-type"].startswith("text/plain")
    _in_workspace(app, "x.svg").write_bytes(b"<svg onload='alert(1)'></svg>")
    r = client.get("/preview/raw", params={"path": "workspace/x.svg"}, headers=auth)
    assert r.headers["content-type"].startswith("text/plain")


def test_raw_with_filename_suffix(client, auth, app):
    """`/preview/raw/<name>` 只是給瀏覽器看的檔名，不參與路徑解析（也不能拿來繞白名單）。"""
    make_png(_in_workspace(app, "shot.png"))
    r = client.get("/preview/raw/shot.png", params={"path": "workspace/shot.png"}, headers=auth)
    assert r.status_code == 200 and r.headers["content-type"].startswith("image/png")
    # 名字亂給也不影響（解析只看 path）
    r = client.get("/preview/raw/anything.txt", params={"path": "workspace/shot.png"}, headers=auth)
    assert r.status_code == 200 and r.headers["content-type"].startswith("image/png")
    r = client.get("/preview/raw/shot.png", params={"path": "/etc/passwd"}, headers=auth)
    assert r.status_code == 404


def test_raw_path_whitelist(client, auth, tmp_path):
    secret = tmp_path / "s.txt"
    secret.write_text("x")
    assert client.get("/preview/raw", params={"path": str(secret)}, headers=auth).status_code == 404
