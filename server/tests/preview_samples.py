"""產生預覽測試用的小樣本檔（docx/xlsx/pptx/pdf/csv/png/md/code）。

測試與真機 QA 共用同一組產生器，避免「測試過但真機長不一樣」。
PDF 用手寫最小 PDF 語法產生，不引入 reportlab。
"""
from __future__ import annotations

import base64
import io
from pathlib import Path

PNG_GRADIENT = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAEAAAAAwCAIAAAAuKetIAAAAWElEQVR42u3TwQ0AEAAEwRM0qf9C9IAPRnw9LrNKMlqyfPvG20PPa64+DwxoBAh8LiAhAgQMkBABAgZIiAABCREg4BNLiAABAyREgIABEiJAQEIECBC46Ew1WQK324EngwAAAABJRU5ErkJggg=="
)
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


# ------------------------------------------------------------------- docx

def make_docx(path: Path) -> Path:
    import docx
    from docx.shared import Inches

    d = docx.Document()
    d.add_heading("季度營運報告", level=1)
    d.add_heading("一、摘要", level=2)
    p = d.add_paragraph()
    p.add_run("本季營收 ").bold = False
    p.add_run("成長 30%").bold = True
    p.add_run("，主因是")
    p.add_run("循環箱").italic = True
    p.add_run(" 導入。")
    d.add_paragraph("第一點", style="List Bullet")
    d.add_paragraph("第二點", style="List Bullet")
    d.add_paragraph("步驟一", style="List Number")
    d.add_heading("二、明細", level=2)
    t = d.add_table(rows=3, cols=3)
    headers = ["品項", "數量", "金額"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    for i, row in enumerate([["循環箱", "1200", "36000"], ["提袋", "800", "9600"]], start=1):
        for j, v in enumerate(row):
            t.rows[i].cells[j].text = v
    d.add_heading("三、附圖", level=2)
    d.add_picture(io.BytesIO(PNG_GRADIENT), width=Inches(1))
    d.save(str(path))
    return path


# ------------------------------------------------------------------- xlsx

def make_xlsx(path: Path, big: bool = False) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "營收"
    ws["A1"] = "月份"
    ws["B1"] = "營收"
    ws["C1"] = "成長率"
    ws["D1"] = "日期"
    ws["E1"] = "合計"
    ws.merge_cells("A1:A1")
    rows = [("一月", 1234567.5, 0.128), ("二月", 2345678.25, 0.301), ("三月", 987654.0, -0.052)]
    import datetime as dt

    for i, (m, rev, g) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=m)
        c = ws.cell(row=i, column=2, value=rev)
        c.number_format = "#,##0.00"
        c2 = ws.cell(row=i, column=3, value=g)
        c2.number_format = "0.0%"
        c3 = ws.cell(row=i, column=4, value=dt.datetime(2026, i - 1, 15))
        c3.number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=5, value=f"=B{i}*2")
    ws.freeze_panes = "B2"
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=3)
    ws.cell(row=5, column=1, value="合併儲存格示範")
    ws2 = wb.create_sheet("備註")
    ws2["A1"] = "說明"
    ws2["A2"] = "第二個工作表"
    if big:
        ws3 = wb.create_sheet("大表")
        for r in range(1, 320):
            for c in range(1, 60):
                ws3.cell(row=r, column=c, value=r * c)
    wb.save(str(path))
    return path


# ------------------------------------------------------------------- pptx

def make_pptx(path: Path, slides: int = 3) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i in range(1, slides + 1):
        layout = prs.slide_layouts[1]
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = f"第 {i} 頁：主題"
        body = s.placeholders[1].text_frame
        body.text = f"重點 {i}-1"
        body.add_paragraph().text = f"重點 {i}-2"
        if i == 2:
            s.shapes.add_picture(io.BytesIO(PNG_GRADIENT), Inches(1), Inches(3), Inches(1), Inches(1))
        s.notes_slide.notes_text_frame.text = f"備註 {i}：這頁要講三十秒"
    prs.save(str(path))
    return path


# -------------------------------------------------------------------- pdf

def make_pdf(path: Path, pages: list[str] | None = None) -> Path:
    """最小可用 PDF：每頁一行 Helvetica 文字。"""
    pages = pages or ["Hello preview page one", "Second page content"]
    objs: list[bytes] = []

    def add(body: bytes) -> int:
        objs.append(body)
        return len(objs)

    font_id = None
    page_ids: list[int] = []
    content_ids: list[int] = []
    # 先佔位：1 catalog, 2 pages
    objs.append(b"")  # 1
    objs.append(b"")  # 2
    font_id = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for text in pages:
        esc = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        stream = f"BT /F1 24 Tf 72 720 Td ({esc}) Tj ET".encode("latin-1", "replace")
        cid = add(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
        content_ids.append(cid)
        pid = add(b"")  # 佔位，等下填（需要知道自己的 id）
        page_ids.append(pid)
    for pid, cid in zip(page_ids, content_ids):
        objs[pid - 1] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 "
            + str(font_id).encode() + b" 0 R >> >> /Contents " + str(cid).encode() + b" 0 R >>"
        )
    kids = b" ".join(str(p).encode() + b" 0 R" for p in page_ids)
    objs[1] = b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(len(page_ids)).encode() + b" >>"
    objs[0] = b"<< /Type /Catalog /Pages 2 0 R >>"

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += str(i).encode() + b" 0 obj\n" + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n"
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += f"{off:010d} 00000 n \n".encode()
    out += (b"trailer\n<< /Size " + str(len(objs) + 1).encode() + b" /Root 1 0 R >>\nstartxref\n"
            + str(xref).encode() + b"\n%%EOF\n")
    path.write_bytes(bytes(out))
    return path


# -------------------------------------------------------------- csv / text

BIG5_CSV_ROWS = [["品項", "數量", "單價"], ["循環箱", "1200", "30"], ["提袋", "800", "12"]]


def make_csv_big5(path: Path) -> Path:
    body = "\n".join(",".join(r) for r in BIG5_CSV_ROWS) + "\n"
    path.write_bytes(body.encode("big5"))
    return path


def make_csv_utf8_bom(path: Path) -> Path:
    body = "name,score\n王小明,90\n李小華,85\n"
    path.write_bytes(b"\xef\xbb\xbf" + body.encode("utf-8"))
    return path


def make_tsv(path: Path) -> Path:
    path.write_text("a\tb\tc\n1\t2\t3\n4\t5\t6\n", encoding="utf-8")
    return path


MD_SAMPLE = """# 循環箱專案報告

一段引言，說明這份文件在講什麼。

## 一、背景

包裝廢棄物在台灣年成長 8%，其中 60% 來自電商一次性紙箱。

- [x] 已完成的項目
- [ ] 還沒完成的項目

## 二、數據

| 指標 | 2025 | 2026 | 變化 |
| --- | ---: | ---: | ---: |
| 循環率 | 10% | 30% | +20pp |
| 成本 | 42 | 28 | -33% |

### 2.1 計算方式

```python
def saving(base: float, rate: float) -> float:
    return base * (1 - rate)
```

## 三、結論

先做 B 端，再做 C 端[^1]。

[^1]: 因為 B 端有固定路線，回收率可控。
"""


def make_markdown(path: Path, repeat: int = 1) -> Path:
    path.write_text(MD_SAMPLE * repeat, encoding="utf-8")
    return path


def make_code(path: Path) -> Path:
    path.write_text(
        "from dataclasses import dataclass\n\n\n"
        "@dataclass\nclass Box:\n    width: int\n    height: int\n\n"
        "    def volume(self, depth: int) -> int:\n        return self.width * self.height * depth\n",
        encoding="utf-8",
    )
    return path


def make_png(path: Path) -> Path:
    path.write_bytes(PNG_GRADIENT)
    return path


def make_all(d: Path) -> dict[str, Path]:
    d.mkdir(parents=True, exist_ok=True)
    return {
        "docx": make_docx(d / "report.docx"),
        "xlsx": make_xlsx(d / "revenue.xlsx", big=True),
        "pptx": make_pptx(d / "deck.pptx"),
        "pdf": make_pdf(d / "two-pages.pdf"),
        "csv_big5": make_csv_big5(d / "big5.csv"),
        "csv_bom": make_csv_utf8_bom(d / "bom.csv"),
        "tsv": make_tsv(d / "data.tsv"),
        "md": make_markdown(d / "report.md", repeat=6),
        "code": make_code(d / "box.py"),
        "png": make_png(d / "pixel.png"),
    }
