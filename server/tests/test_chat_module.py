"""A. 聊天模組：session 管理、分類、搜尋、上傳／下載／預覽、模型清單、Hermes 歷史、WS 附件／重生成／編輯。"""
from __future__ import annotations

import base64
import io
import json
import sqlite3
import time
from pathlib import Path

import pytest

PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _agent(client, auth, profile="default"):
    return next(a["id"] for a in client.get("/agents", headers=auth).json() if a["profile"] == profile)


def _session(client, auth, profile="default", **kw):
    return client.post("/sessions", json={"agent_id": _agent(client, auth, profile), **kw}, headers=auth).json()


def _collect(ws, until=("run.completed", "run.failed", "run.cancelled")):
    events = []
    while True:
        ev = json.loads(ws.receive_text())
        events.append(ev)
        if ev["type"] in until:
            return events


# ---------------------------------------------------------------- sessions

def test_session_patch_rename_archive_model_and_ordering(client, auth):
    a = _session(client, auth, title="A")
    b = _session(client, auth, title="B")
    r = client.patch(f"/sessions/{a['id']}", json={"title": "  改名  "}, headers=auth)
    assert r.status_code == 200 and r.json()["title"] == "改名"
    assert client.patch(f"/sessions/{a['id']}", json={"title": "   "}, headers=auth).status_code == 400
    # model per session
    r = client.post(f"/sessions/{a['id']}/model", json={"model": "anthropic/claude-sonnet-5", "provider": "nous"}, headers=auth)
    assert r.json()["model"] == "anthropic/claude-sonnet-5" and r.json()["provider"] == "nous"
    assert client.get(f"/sessions/{a['id']}", headers=auth).json()["usage"]["total_tokens"] == 0
    # archive hides from default list
    client.patch(f"/sessions/{b['id']}", json={"archived": True}, headers=auth)
    ids = [s["id"] for s in client.get("/sessions", headers=auth).json()]
    assert ids == [a["id"]]
    ids = [s["id"] for s in client.get("/sessions", params={"include_archived": "true"}, headers=auth).json()]
    assert set(ids) == {a["id"], b["id"]}
    assert client.patch(f"/sessions/{b['id']}", json={"archived": False}, headers=auth).json()["archived"] is False


def test_session_list_running_first_then_recent(client, auth, token):
    old = _session(client, auth, title="old")
    time.sleep(0.01)
    new = _session(client, auth, title="new")
    ids = [s["id"] for s in client.get("/sessions", headers=auth).json()]
    assert ids == [new["id"], old["id"]]
    # a run on `old` moves it up (last_message_at) and marks completed
    with client.websocket_connect(f"/ws/chat?token={token}") as ws:
        ws.receive_text()
        ws.send_text(json.dumps({"type": "run", "session_id": old["id"], "input": "hi"}))
        _collect(ws)
    rows = client.get("/sessions", headers=auth).json()
    assert rows[0]["id"] == old["id"] and rows[0]["run_status"] == "completed" and rows[0]["running"] is False
    assert rows[0]["usage"]["total_tokens"] == 3 and rows[0]["usage"]["context_tokens"] == 3


def test_categories_crud_and_assign(client, auth):
    r = client.post("/chat/categories", json={"name": "行銷", "color": "#f59e0b"}, headers=auth)
    assert r.status_code == 201
    cat = r.json()
    assert client.post("/chat/categories", json={"name": " "}, headers=auth).status_code == 400
    s = _session(client, auth, category_id=cat["id"])
    assert s["category_id"] == cat["id"]
    assert client.patch(f"/chat/categories/{cat['id']}", json={"name": "行銷2"}, headers=auth).json()["name"] == "行銷2"
    assert [c["name"] for c in client.get("/chat/categories", headers=auth).json()] == ["行銷2"]
    assert [x["id"] for x in client.get("/sessions", params={"category_id": cat["id"]}, headers=auth).json()] == [s["id"]]
    assert client.patch(f"/sessions/{s['id']}", json={"category_id": "cat_nope"}, headers=auth).status_code == 404
    assert client.patch(f"/sessions/{s['id']}", json={"category_id": ""}, headers=auth).json()["category_id"] is None
    client.patch(f"/sessions/{s['id']}", json={"category_id": cat["id"]}, headers=auth)
    assert client.delete(f"/chat/categories/{cat['id']}", headers=auth).json()["ok"]
    assert client.get(f"/sessions/{s['id']}", headers=auth).json()["category_id"] is None
    # other member cannot see my categories
    client.post("/members", json={"username": "eve", "password": "evepass", "role": "member"}, headers=auth)
    eve = {"Authorization": "Bearer " + client.post("/auth/login", json={"username": "eve", "password": "evepass"}).json()["token"]}
    client.post("/chat/categories", json={"name": "mine"}, headers=auth)
    assert client.get("/chat/categories", headers=eve).json() == []


def test_search_titles_and_messages(client, auth, token):
    s1 = _session(client, auth, title="季度報告討論")
    s2 = _session(client, auth, title="其他")
    with client.websocket_connect(f"/ws/chat?token={token}") as ws:
        ws.receive_text()
        ws.send_text(json.dumps({"type": "run", "session_id": s2["id"], "input": "請整理報告的重點"}))
        _collect(ws)
    hits = client.get("/sessions/search", params={"q": "報告"}, headers=auth).json()
    by_id = {h["session"]["id"]: h for h in hits}
    assert by_id[s1["id"]]["match"] == "title"
    assert by_id[s2["id"]]["match"] == "message" and "報告" in by_id[s2["id"]]["snippet"]
    assert client.get("/sessions/search", params={"q": ""}, headers=auth).json() == []
    assert client.get("/sessions/search", params={"q": "zzz-nothing"}, headers=auth).json() == []


# ---------------------------------------------------------------- uploads / files / preview

def test_upload_download_preview_and_whitelist(client, auth, tmp_path, app):
    s = _session(client, auth)
    files = [("files", ("報 告.md", b"# Title\n\nhello *md*", "text/markdown")),
             ("files", ("pic.png", PNG_1PX, "image/png")),
             ("files", ("../../evil.txt", b"x", "text/plain"))]
    r = client.post("/chat/uploads", data={"session_id": s["id"]}, files=files, headers=auth)
    assert r.status_code == 201, r.text
    up = r.json()
    names = [u["name"] for u in up]
    assert names == ["報 告.md", "pic.png", "evil.txt"]  # path traversal stripped
    root = Path(app.state.settings.db_path).parent / "uploads"
    for u in up:
        assert Path(u["path"]).is_file() and Path(u["path"]).resolve().is_relative_to(root.resolve())
    assert len(client.get("/chat/uploads", params={"session_id": s["id"]}, headers=auth).json()) == 3
    # download (inline + attachment)
    r = client.get("/chat/files", params={"path": up[1]["path"]}, headers=auth)
    assert r.status_code == 200 and r.headers["content-type"].startswith("image/png") and r.content == PNG_1PX
    r = client.get("/chat/files", params={"path": up[0]["path"], "download": "true"}, headers=auth)
    assert "attachment" in r.headers["content-disposition"]
    # preview markdown
    pv = client.get("/chat/preview", params={"path": up[0]["path"]}, headers=auth).json()
    assert pv["kind"] == "markdown" and pv["text"].startswith("# Title")
    assert client.get("/chat/preview", params={"path": up[1]["path"]}, headers=auth).json()["kind"] == "image"
    # whitelist: arbitrary file outside uploads/workspace/messages → 404
    secret = tmp_path / "secret.txt"
    secret.write_text("nope")
    assert client.get("/chat/files", params={"path": str(secret)}, headers=auth).status_code == 404
    assert client.get("/chat/preview", params={"path": str(secret)}, headers=auth).status_code == 404
    assert client.get("/chat/files", params={"path": "relative/x"}, headers=auth).status_code == 404
    # other member cannot read my upload
    client.post("/members", json={"username": "eve", "password": "evepass", "role": "member"}, headers=auth)
    eve = {"Authorization": "Bearer " + client.post("/auth/login", json={"username": "eve", "password": "evepass"}).json()["token"]}
    # (same company uploads root is readable by company members — documented; but session upload requires ownership)
    assert client.post("/chat/uploads", data={"session_id": s["id"]}, files=files[:1], headers=eve).status_code == 404


def test_download_path_mentioned_in_message_is_allowed(client, auth, token, tmp_path, gw_state):
    out = tmp_path / "agent-output.csv"
    out.write_text("a,b\n1,2\n3,4\n")
    s = _session(client, auth)
    # agent "produces" the file: its path shows up in the assistant output
    with client.websocket_connect(f"/ws/chat?token={token}") as ws:
        ws.receive_text()
        ws.send_text(json.dumps({"type": "run", "session_id": s["id"], "input": f"存到 {out}"}))
        _collect(ws)
    pv = client.get("/chat/preview", params={"path": str(out)}, headers=auth).json()
    assert pv["kind"] == "csv" and pv["rows"] == [["a", "b"], ["1", "2"], ["3", "4"]]
    assert client.get("/chat/files", params={"path": str(out)}, headers=auth).status_code == 200
    # a sibling file never mentioned stays hidden
    (tmp_path / "other.csv").write_text("x")
    assert client.get("/chat/files", params={"path": str(tmp_path / 'other.csv')}, headers=auth).status_code == 404


def test_office_previews(client, auth):
    import docx
    import openpyxl
    from pptx import Presentation

    s = _session(client, auth)
    d = docx.Document()
    d.add_heading("標題一", level=1)
    d.add_paragraph("內文段落 <b>")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "c1", "c2"
    bd = io.BytesIO(); d.save(bd)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "S1"; ws.append(["h1", "h2"]); ws.append([1, 2.5])
    bx = io.BytesIO(); wb.save(bx)
    prs = Presentation(); sl = prs.slides.add_slide(prs.slide_layouts[1]); sl.shapes.title.text = "投影片標題"
    sl.placeholders[1].text = "要點"
    bp = io.BytesIO(); prs.save(bp)
    files = [("files", ("a.docx", bd.getvalue())), ("files", ("b.xlsx", bx.getvalue())), ("files", ("c.pptx", bp.getvalue()))]
    up = client.post("/chat/uploads", data={"session_id": s["id"]}, files=files, headers=auth).json()
    pv = client.get("/chat/preview", params={"path": up[0]["path"]}, headers=auth).json()
    assert pv["kind"] == "docx" and "<h1>標題一</h1>" in pv["html"] and "&lt;b&gt;" in pv["html"] and "<td>c2</td>" in pv["html"]
    pv = client.get("/chat/preview", params={"path": up[1]["path"]}, headers=auth).json()
    assert pv["kind"] == "xlsx" and pv["sheets"][0]["name"] == "S1" and pv["sheets"][0]["rows"] == [["h1", "h2"], [1, 2.5]]
    pv = client.get("/chat/preview", params={"path": up[2]["path"]}, headers=auth).json()
    assert pv["kind"] == "pptx" and "投影片標題" in pv["html"] and "要點" in pv["html"]


def test_html_upload_is_served_as_text(client, auth):
    s = _session(client, auth)
    up = client.post("/chat/uploads", data={"session_id": s["id"]}, files=[("files", ("x.html", b"<script>1</script>"))], headers=auth).json()
    r = client.get("/chat/files", params={"path": up[0]["path"]}, headers=auth)
    assert r.headers["content-type"].startswith("text/plain")
    assert client.get("/chat/preview", params={"path": up[0]["path"]}, headers=auth).json()["kind"] == "html"


# ---------------------------------------------------------------- models

def test_models_fallback_to_v1_models_and_normalize(client, auth):
    r = client.get("/chat/models", params={"profile": "researcher"}, headers=auth).json()
    assert r["fallback"] is True and [m["id"] for m in r["models"]] == ["hermes-agent"]
    from studio.modules.chat.models_api import normalize_options
    n = normalize_options({"model": "x/y", "provider": "nous", "providers": [
        {"slug": "nous", "name": "Nous", "is_current": True, "models": ["x/y", "x/z"], "pricing": {"x/y": {"input": "$1", "output": "$2", "free": False}}},
        {"slug": "openai", "models": [{"id": "gpt"}]}]})
    assert [m["id"] for m in n["models"]] == ["x/y", "x/z", "gpt"]
    assert n["models"][0]["pricing"]["input"] == "$1" and n["models"][2]["provider"] == "openai"
    assert n["current"] == {"model": "x/y", "provider": "nous"}


# ---------------------------------------------------------------- hermes history

def _make_state_db(path: Path, sid: str, source: str, title: str):
    conn = sqlite3.connect(path)
    conn.executescript("""
    CREATE TABLE sessions (id TEXT PRIMARY KEY, source TEXT NOT NULL, user_id TEXT, model TEXT, started_at REAL NOT NULL,
      ended_at REAL, message_count INTEGER DEFAULT 0, tool_call_count INTEGER DEFAULT 0, input_tokens INTEGER DEFAULT 0,
      output_tokens INTEGER DEFAULT 0, title TEXT, display_name TEXT, archived INTEGER DEFAULT 0, last_activity_at REAL);
    CREATE TABLE messages (id INTEGER PRIMARY KEY AUTOINCREMENT, session_id TEXT NOT NULL, role TEXT NOT NULL, content TEXT,
      tool_call_id TEXT, tool_calls TEXT, tool_name TEXT, timestamp REAL NOT NULL, reasoning TEXT);
    """)
    t0 = 1_787_900_000.0
    conn.execute("INSERT INTO sessions (id, source, model, started_at, message_count, input_tokens, output_tokens, title) VALUES (?,?,?,?,?,?,?,?)",
                 (sid, source, "gpt-x", t0, 4, 100, 50, title))
    conn.execute("INSERT INTO sessions (id, source, started_at, message_count) VALUES ('empty', 'cli', ?, 0)", (t0 - 10,))
    rows = [
        (sid, "system", "sys prompt", None, None, None, t0),
        (sid, "user", "查一下天氣", None, None, None, t0 + 1),
        (sid, "assistant", "", None, json.dumps([{"id": "call_1", "type": "function", "function": {"name": "web_search", "arguments": "{\"q\": \"weather\"}"}}]), None, t0 + 2),
        (sid, "tool", "{\"temp\": 30}", "call_1", None, "web_search", t0 + 3),
        (sid, "assistant", "今天 30 度", None, None, None, t0 + 4),
    ]
    conn.executemany("INSERT INTO messages (session_id, role, content, tool_call_id, tool_calls, tool_name, timestamp) VALUES (?,?,?,?,?,?,?)", rows)
    conn.commit()
    conn.close()


def test_hermes_history_list_messages_import(client, auth, hermes_home):
    _make_state_db(hermes_home / "state.db", "tg_1", "telegram", "天氣對話")
    (hermes_home / "profiles" / "researcher" / "state.db").touch()
    _make_state_db(hermes_home / "profiles" / "researcher" / "state.db", "cli_1", "cli", "researcher 對話")
    src = client.get("/chat/hermes-history/sources", headers=auth).json()
    assert {x["profile"]: x["sources"] for x in src} == {"default": {"telegram": 1}, "researcher": {"cli": 1}}
    lst = client.get("/chat/hermes-history", headers=auth).json()
    assert [(x["profile"], x["id"], x["source"]) for x in lst] == [("default", "tg_1", "telegram"), ("researcher", "cli_1", "cli")]
    assert lst[0]["title"] == "天氣對話" and lst[0]["message_count"] == 4
    assert [x["id"] for x in client.get("/chat/hermes-history", params={"profile": "researcher"}, headers=auth).json()] == ["cli_1"]
    assert client.get("/chat/hermes-history", params={"q": "researcher"}, headers=auth).json()[0]["id"] == "cli_1"
    m = client.get("/chat/hermes-history/default/tg_1/messages", headers=auth).json()
    assert [(x["role"], x.get("tool_name")) for x in m["messages"]] == [("user", None), ("tool", "web_search"), ("assistant", None)]
    assert m["messages"][1]["tool_args"] == {"q": "weather"} and m["messages"][1]["tool_result"] == '{"temp": 30}'
    assert client.get("/chat/hermes-history/nope/tg_1/messages", headers=auth).status_code == 404
    assert client.get("/chat/hermes-history/default/nope/messages", headers=auth).status_code == 404
    # import → Studio session bound to agent with that profile
    r = client.post("/chat/hermes-history/researcher/cli_1/import", json={}, headers=auth)
    assert r.status_code == 201, r.text
    s = r.json()
    researcher = next(a for a in client.get("/agents", headers=auth).json() if a["profile"] == "researcher")
    assert s["agent_id"] == researcher["id"] and s["source"] == "cli" and s["imported_from"] == "researcher:cli_1" and s["title"] == "researcher 對話"
    assert s["usage"]["total_tokens"] == 150 and s["model"] == "gpt-x"
    msgs = client.get(f"/sessions/{s['id']}/messages", headers=auth).json()
    assert [x["role"] for x in msgs] == ["user", "tool", "assistant"] and msgs[1]["tool_args"] == {"q": "weather"}
    # idempotent
    assert client.post("/chat/hermes-history/researcher/cli_1/import", json={}, headers=auth).json()["id"] == s["id"]
    # state.db untouched (still read-only): message count unchanged
    conn = sqlite3.connect(hermes_home / "profiles" / "researcher" / "state.db")
    assert conn.execute("select count(*) from messages").fetchone()[0] == 5
    conn.close()


# ---------------------------------------------------------------- WS additions

def test_ws_attachments_reply_and_image_parts(client, auth, token, gw_state, hermes_home):
    s = _session(client, auth)
    up = client.post("/chat/uploads", data={"session_id": s["id"]},
                     files=[("files", ("pic.png", PNG_1PX, "image/png")), ("files", ("note.txt", b"hi", "text/plain"))],
                     headers=auth).json()
    with client.websocket_connect(f"/ws/chat?token={token}") as ws:
        ws.receive_text()
        ws.send_text(json.dumps({"type": "run", "session_id": s["id"], "input": "看圖", "attachments": up}))
        ev = _collect(ws)
        started = ev[0]
        assert started["type"] == "run.started" and [a["name"] for a in started["attachments"]] == ["pic.png", "note.txt"]
        assert ev[-1]["type"] == "run.completed" and ev[-1]["session_usage"]["total_tokens"] == 3 and ev[-1]["message_id"]
        body = gw_state.runs[started["run_id"]]["body"]
        # image → content parts on a message list; paths still appended to the text
        assert isinstance(body["input"], list) and body["input"][0]["role"] == "user"
        parts = body["input"][0]["content"]
        assert parts[0]["type"] == "text" and up[1]["path"] in parts[0]["text"] and "看圖" in parts[0]["text"]
        assert parts[1]["type"] == "image_url" and parts[1]["image_url"]["url"].startswith("data:image/png;base64,")
        # reply_to quotes the assistant message into the input
        msgs = client.get(f"/sessions/{s['id']}/messages", headers=auth).json()
        assert msgs[0]["attachments"][0]["name"] == "pic.png" and msgs[-1]["usage"]["total_tokens"] == 3
        ws.send_text(json.dumps({"type": "run", "session_id": s["id"], "input": "為什麼", "reply_to": msgs[-1]["id"]}))
        ev2 = _collect(ws)
        body2 = gw_state.runs[ev2[0]["run_id"]]["body"]
        assert body2["input"].startswith("[引用先前訊息]\n> ") and body2["input"].endswith("為什麼")
        assert ev2[0]["reply_to"] == msgs[-1]["id"]
    msgs = client.get(f"/sessions/{s['id']}/messages", headers=auth).json()
    assert msgs[2]["reply_to"] == msgs[1]["id"] and msgs[2]["content"] == "為什麼"  # stored content is the raw user text


def test_ws_regenerate_and_edit(client, auth, token, gw_state):
    s = _session(client, auth)
    with client.websocket_connect(f"/ws/chat?token={token}") as ws:
        ws.receive_text()
        ws.send_text(json.dumps({"type": "run", "session_id": s["id"], "input": "第一"}))
        _collect(ws)
        ws.send_text(json.dumps({"type": "run", "session_id": s["id"], "input": "第二"}))
        _collect(ws)
        before = client.get(f"/sessions/{s['id']}/messages", headers=auth).json()
        assert [m["content"] for m in before] == ["第一", "echo[default|h=0]: 第一", "第二", "echo[default|h=2]: 第二"]
        # regenerate: removes last user+assistant, re-runs "第二"
        ws.send_text(json.dumps({"type": "regenerate", "session_id": s["id"]}))
        removed = json.loads(ws.receive_text())
        assert removed["type"] == "messages.removed" and set(removed["ids"]) == {before[2]["id"], before[3]["id"]}
        ev = _collect(ws)
        assert ev[-1]["output"] == "echo[default|h=2]: 第二"
        after = client.get(f"/sessions/{s['id']}/messages", headers=auth).json()
        assert [m["content"] for m in after] == ["第一", "echo[default|h=0]: 第一", "第二", "echo[default|h=2]: 第二"]
        assert after[2]["id"] != before[2]["id"]
        # edit last user message
        ws.send_text(json.dumps({"type": "edit", "session_id": s["id"], "input": "第二改"}))
        removed = json.loads(ws.receive_text())
        assert removed["type"] == "messages.removed" and len(removed["ids"]) == 2
        ev = _collect(ws)
        assert ev[-1]["output"] == "echo[default|h=2]: 第二改"
        ws.send_text(json.dumps({"type": "edit", "session_id": s["id"], "input": ""}))
        assert json.loads(ws.receive_text())["code"] == "bad_message"
    final = client.get(f"/sessions/{s['id']}/messages", headers=auth).json()
    assert [m["content"] for m in final] == ["第一", "echo[default|h=0]: 第一", "第二改", "echo[default|h=2]: 第二改"]


def test_ws_session_model_override_reaches_gateway(client, auth, token, gw_state):
    s = _session(client, auth)
    client.post(f"/sessions/{s['id']}/model", json={"model": "openai/gpt-5.5", "provider": "nous"}, headers=auth)
    with client.websocket_connect(f"/ws/chat?token={token}") as ws:
        ws.receive_text()
        ws.send_text(json.dumps({"type": "run", "session_id": s["id"], "input": "x"}))
        ev = _collect(ws)
        body = gw_state.runs[ev[0]["run_id"]]["body"]
        assert body["model"] == "openai/gpt-5.5" and body["provider"] == "nous"
        # one-off override wins
        ws.send_text(json.dumps({"type": "run", "session_id": s["id"], "input": "y", "model": "z/zz"}))
        ev = _collect(ws)
        assert gw_state.runs[ev[0]["run_id"]]["body"]["model"] == "z/zz"


def test_ensure_columns_migrates_old_db(tmp_path):
    from sqlalchemy import text
    from studio.db import make_engine
    from studio.modules.chat import ensure_columns
    engine = make_engine(tmp_path / "old.db")
    with engine.begin() as c:
        c.execute(text("CREATE TABLE sessions (id VARCHAR PRIMARY KEY, title VARCHAR, run_status VARCHAR DEFAULT 'running')"))
        c.execute(text("CREATE TABLE messages (id VARCHAR PRIMARY KEY, content VARCHAR)"))
        c.execute(text("INSERT INTO sessions (id, title, run_status) VALUES ('s1', 't', 'running')"))
    assert ensure_columns(engine) == 13
    with engine.begin() as c:
        cols = {r[1] for r in c.execute(text("PRAGMA table_info(sessions)")).fetchall()}
        assert {"archived", "model", "context_tokens", "imported_from"} <= cols
        assert c.execute(text("SELECT run_status FROM sessions")).fetchone()[0] == ""
    assert ensure_columns(engine) == 0
